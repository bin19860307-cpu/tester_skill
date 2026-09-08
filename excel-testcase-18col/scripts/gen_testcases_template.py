#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
18列企业标准测试用例 Excel 生成脚本模板
模板格式参考：班级管理测试用例_v2.xlsx
18列：实体标识,编号,名称,归属套件,版本,类型,可测试,基础用例,状态,预置条件,测试过程,接收标准,创建人,创建时间,描述,脚本名称,验证话单,信令流程

使用方法：
  1. 修改 OUTPUT_PATH 为实际输出路径
  2. 修改 CREATOR / VERSION 为实际值
  3. 在「用例数据区域」填写 suite() 和 tc() 调用
  4. 运行脚本：python gen_testcases_template.py
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date

# ============================================================
# 配置区域（按需修改）
# ============================================================
OUTPUT_PATH = r"C:\替换为实际输出路径\测试用例_v2.xlsx"
TODAY   = date.today().strftime("%Y-%m-%d")
VERSION = "v1.0"
CREATOR = "QA"

# ============================================================
# 样式定义（对齐班级管理模板风格）
# ============================================================
HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")   # 深蓝
HEADER_FONT  = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=10)
SUITE_FILL   = PatternFill("solid", fgColor="D9E1F2")   # 浅蓝 - 顶层/一级套件
SUITE_FONT   = Font(name="Microsoft YaHei", bold=True, color="1F4E79", size=10)
SUITE2_FILL  = PatternFill("solid", fgColor="EDF2F8")   # 淡蓝 - 二级套件
SUITE2_FONT  = Font(name="Microsoft YaHei", bold=True, color="17375E", size=10)
NORMAL_FONT  = Font(name="Microsoft YaHei", size=9)
THIN   = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP   = Alignment(wrap_text=True, vertical="top")

# ============================================================
# 列定义（18列，顺序与模板一致）
# ============================================================
HEADERS = [
    "实体标识", "编号", "名称", "归属套件",
    "版本", "类型", "可测试", "基础用例",
    "状态", "预置条件", "测试过程", "接收标准",
    "创建人", "创建时间", "描述", "脚本名称",
    "验证话单", "信令流程"
]
# 对应各列宽度（字符数）
COL_WIDTHS = [10, 10, 48, 10, 8, 10, 8, 8, 8, 36, 52, 52, 8, 12, 30, 14, 8, 8]

# ============================================================
# 数据容器
# ============================================================
DATA = []


def suite(code, name, parent, desc=""):
    """
    添加用例套件行。
    
    Args:
        code   (int): 套件编号。顶层=1，一级=101/102...
        name   (str): 套件名称
        parent (int): 父级套件编号。顶层套件填0，一级套件填1
        desc   (str): 套件描述（可选）
    """
    DATA.append({
        "实体标识": "用例套件",
        "编号":    code,
        "名称":    name,
        "归属套件": parent,
        "描述":    desc,
    })


def tc(code, name, parent, pre, steps, accept, tc_type="功能测试", desc=""):
    """
    添加测试用例行。
    
    Args:
        code    (int): 用例编号，5位数，前3位=所属套件编号（如10101）
        name    (str): 用例名称
        parent  (int): 所属套件编号（如101）
        pre     (str): 预置条件
        steps   (str): 测试步骤，用 \\n 分隔各步骤
        accept  (str): 接收标准（预期结果）
        tc_type (str): 测试类型，默认"功能测试"，可选"异常测试"/"边界测试"等
        desc    (str): 用例描述（可选）
    """
    DATA.append({
        "实体标识": "测试用例",
        "编号":    code,
        "名称":    name,
        "归属套件": parent,
        "版本":    VERSION,
        "类型":    tc_type,
        "可测试":  "是",
        "基础用例": "否",
        "状态":    "待评审",
        "预置条件": pre,
        "测试过程": steps,
        "接收标准": accept,
        "创建人":  CREATOR,
        "创建时间": TODAY,
        "描述":    desc,
    })


# ============================================================
# ====== 用例数据区域（在此填写所有 suite() 和 tc() 调用）======
# ============================================================

# 顶层套件（归属套件=0）
suite(1, "项目名称/功能模块", 0, "顶层套件说明")

# 一级套件示例（归属套件=1）
suite(101, "功能模块A", 1, "模块A测试范围说明")

# 测试用例示例
tc(10101, "正向流程-基础功能验证", 101,
   "已登录系统，进入目标页面",
   "1. 登录系统\n2. 进入目标页面\n3. 执行操作步骤\n4. 观察结果",
   "操作成功，页面展示预期内容，无报错")

tc(10102, "异常场景-非法输入处理", 101,
   "已登录系统，进入目标页面",
   "1. 进入目标页面\n2. 输入非法数据\n3. 提交",
   "系统给出明确错误提示，不允许提交，不产生脏数据",
   tc_type="异常测试")

# 继续添加更多套件和用例...


# ============================================================
# 写入 Excel
# ============================================================
def create_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "测试用例表"

    # 写入标题行
    for ci, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = BORDER
    ws.row_dimensions[1].height = 28

    # 设置列宽
    for ci, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    # 写入数据行
    row = 2
    suite_cnt = tc_cnt = 0
    for record in DATA:
        is_suite = (record.get("实体标识") == "用例套件")
        parent   = record.get("归属套件", 0)

        # 判断套件层级
        if is_suite:
            if isinstance(parent, int) and parent > 1:
                fill, font = SUITE2_FILL, SUITE2_FONT  # 二级套件（淡蓝）
            else:
                fill, font = SUITE_FILL, SUITE_FONT     # 顶层/一级套件（浅蓝）
            suite_cnt += 1
        else:
            tc_cnt += 1

        # 写入各列
        for ci, h in enumerate(HEADERS, 1):
            cell = ws.cell(row=row, column=ci, value=record.get(h, ""))
            cell.border = BORDER
            if is_suite:
                cell.fill      = fill
                cell.font      = font
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                cell.font      = NORMAL_FONT
                cell.alignment = WRAP

        # 自适应行高
        if is_suite:
            ws.row_dimensions[row].height = 18
        else:
            lines = max(
                str(record.get("测试过程", "")).count("\n") + 1,
                str(record.get("接收标准", "")).count("\n") + 1,
            )
            ws.row_dimensions[row].height = max(lines * 14 + 6, 20)

        row += 1

    ws.freeze_panes = "A2"
    wb.save(OUTPUT_PATH)
    print(f"✅ 测试用例已生成：{OUTPUT_PATH}")
    print(f"📊 套件数：{suite_cnt}  测试用例数：{tc_cnt}")


if __name__ == "__main__":
    create_workbook()
